import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_word_doc():
    doc = Document()
    
    # Header Information
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("Subject: MACHINE LEARNING\nName: RAAGHAV AGARWAL\nRegistration Number: 24BAI1453")
    run.font.size = Pt(11)
    run.bold = True
    
    doc.add_paragraph("").add_run().add_break()
    
    # Title
    title = doc.add_heading("Project Report: Parkinson's Voice Companion", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("").add_run().add_break()
    
    # 1. Table of Contents
    doc.add_heading("1. Suggested Document Topics (Table of Contents)", level=1)
    doc.add_paragraph("If you are submitting this as a final year project or research report, structure your document using these topics:")
    
    toc = [
        "1. Abstract / Executive Summary",
        "2. Introduction\n   - Problem Statement: Challenges in continuous, non-invasive Parkinson's disease monitoring.\n   - Objectives: Utilizing acoustic biomarkers and machine learning for early detection and disease tracking.",
        "3. Literature Review\n   - Clinical relevance of Dysphonia in Parkinson's (Little et al., Tsanas et al.)\n   - LSVT LOUD® therapy efficacy.",
        "4. Methodology & Data Architecture\n   - Data Acquisition (Datasets Used)\n   - Digital Signal Processing (DSP) & Feature Extraction\n   - Multimodal Integration (Voice + Motor Tapping)",
        "5. System Architecture\n   - Backend (FastAPI, Librosa, Scikit-Learn)\n   - Frontend (Web Audio API, Three.js, Chart.js)",
        "6. Machine Learning Models\n   - Classification Pipeline (SVM, Random Forest, CNN-BiLSTM)\n   - Regression Pipeline (UPDRS Telemonitoring)\n   - Anomaly Detection (Autoencoder, BOCPD)",
        "7. Implementation Details (What We Have Done)",
        "8. Results & UI Screenshots",
        "9. Conclusion & Future Scope"
    ]
    for item in toc:
        doc.add_paragraph(item, style='List Bullet')
        
    doc.add_page_break()

    # 2. What We Have Done
    doc.add_heading("2. What We Have Done (Project Implementation Summary)", level=1)
    doc.add_paragraph("We successfully engineered a full-stack, ML-powered clinical web application designed for non-invasive Parkinson's Disease (PD) monitoring, early detection, and vocal therapy.")
    doc.add_heading("Key Achievements:", level=2)
    
    achievements = [
        "Real-time Acoustic Biomarker Extraction Engine: We built a custom Digital Signal Processing (DSP) pipeline using Python (librosa, parselmouth) that ingests a patient's sustained vowel phonation (/a/) and deterministically calculates 22 clinical acoustic biomarkers (including Jitter, Shimmer, NHR, HNR, Pitch Period Entropy (PPE), RPDE, and DFA).",
        "Ensemble Machine Learning Diagnostic Assessment: We trained and deployed a multi-model consensus system directly into the backend. By passing extracted biomarkers through an SVM, Random Forest, a CNN-BiLSTM deep learning network, and an Autoencoder, the system provides a highly accurate, data-driven probability of Parkinsonian dysphonia.",
        "Longitudinal Telemonitoring System: We implemented a system to predict MDS-UPDRS (Unified Parkinson's Disease Rating Scale) motor and total severity scores solely from voice. We layered Bayesian Online Changepoint Detection (BOCPD) on top of this to detect statistically significant medication state shifts (e.g., Levodopa \"On/Off\" phases).",
        "Interactive LSVT LOUD® Therapy Coach: We created a browser-based biofeedback module that captures real-time microphone audio to evaluate sustained phonation duration (MPT) and vocal loudness (dB SPL), providing a visual gauge for patients practicing the clinical LSVT LOUD® protocol.",
        "Multimodal Motor Fusion Module: We integrated a browser-based finger-tapping test (simulating MDS-UPDRS Item 3.4) to assess tap frequency, rhythmicity (CV), and fatigue. This motor data is mathematically fused with the voice classification probabilities to generate a holistic, multi-modal composite risk score.",
        "Premium 3D Clinical Interface: We designed a visually stunning, premium frontend using glassmorphism, Three.js 3D neural particle shaders, interactive Chart.js radar maps, and custom 3D generated anatomical hero images, moving away from \"generic AI\" aesthetics to a professional biotech aesthetic."
    ]
    for ach in achievements:
        p = doc.add_paragraph(style='List Number')
        split_ach = ach.split(":", 1)
        if len(split_ach) == 2:
            p.add_run(split_ach[0] + ":").bold = True
            p.add_run(split_ach[1])
        else:
            p.add_run(ach)

    doc.add_page_break()

    # 3. Screenshots
    doc.add_heading("3. Results & UI Screenshots", level=1)
    
    doc.add_heading("Figure 1: Voice Biomarker Lab & Diagnostic Dashboard", level=2)
    img_path = r"C:\Users\Raaghav Agarwal\.gemini\antigravity-ide\brain\57f9f74e-84e8-4d77-87d1-93aa9a1ad697\warm_palette_v3_1791310115504.png"
    if os.path.exists(img_path):
        doc.add_picture(img_path, width=Inches(6.0))
    p1 = doc.add_paragraph("Caption: The main acoustic diagnostic dashboard highlighting the real-time extraction of 22 biomarkers mapped against normative healthy baselines.")
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph().add_run().add_break()
    
    doc.add_heading("Figure 2: Longitudinal Telemonitoring & Disease Progression", level=2)
    p2 = doc.add_paragraph("Caption: Longitudinal tracking of predicted UPDRS scores and acoustic entropy, featuring automatic Changepoint Detection for medication state monitoring.")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading("Figure 3: LSVT LOUD® Therapy Biofeedback Interface", level=2)
    p3 = doc.add_paragraph("Caption: Real-time visual biofeedback interface designed to coach patients through sustained high-intensity phonation exercises.")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading("Figure 4: Multimodal Motor Fluctuation (Finger Tapping) Test", level=2)
    p4 = doc.add_paragraph("Caption: Browser-based MDS-UPDRS Item 3.4 finger-tapping simulation integrated with voice metrics for a multimodal composite severity score.")
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # 4. Dataset Frame Architecture
    doc.add_heading("4. Dataset Frame Architecture", level=1)
    
    # Diagram Placeholder
    p_ph = doc.add_paragraph()
    p_ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ph = p_ph.add_run("\n\n[ INSERT DATA ARCHITECTURE FLOWCHART SCREENSHOT HERE ]\n\n")
    run_ph.bold = True
    run_ph.font.size = Pt(14)
    run_ph.font.color.rgb = docx.shared.RGBColor(255, 0, 0) if 'docx.shared' in globals() else None

    pdia = doc.add_paragraph("Caption: Data Architecture Flowchart illustrating ingestion, deterministic DSP engine, and parallel ML inference pipeline.")
    pdia.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().add_run().add_break()

    doc.add_paragraph("To ground the application in validated clinical science, our ML models and normative baselines were established using three foundational, peer-reviewed public datasets.")
    
    doc.add_heading("A. Datasets Used", level=2)
    
    datasets = [
        "1. UCI Oxford Parkinson's Disease Acoustic Dataset (Little et al., 2008)\n   - Scale: 31 subjects (23 with PD, 8 healthy controls), 195 voice recordings.\n   - Role: Used to establish the foundational normative baselines for healthy controls and to train the primary SVM and Random Forest classification models.\n   - Target: Binary Classification (Healthy vs. PD).",
        "2. UCI Parkinson's Telemonitoring Dataset (Tsanas et al., 2010)\n   - Scale: 42 early-stage PD subjects, 5,875 voice recordings gathered over 6 months in home environments.\n   - Role: Used to train the regression models that map specific acoustic perturbations to clinical MDS-UPDRS motor and total severity scores.\n   - Target: Continuous Regression (UPDRS estimation).",
        "3. Sakar et al. Dataset (2019)\n   - Scale: 252 subjects, 756 recordings. Includes advanced baseline features (MFCCs, Wavelet Transforms, Vocal Folds features).\n   - Role: Used to supplement the training of our more complex Deep Learning sequence model (CNN-BiLSTM) to improve generalizability across broader acoustic profiles."
    ]
    for ds in datasets:
        p = doc.add_paragraph(style='List Bullet')
        split_ds = ds.split("\n", 1)
        p.add_run(split_ds[0]).bold = True
        if len(split_ds) > 1:
            p.add_run("\n" + split_ds[1])

    doc.add_heading("B. Architectural Flow (How Data Moves)", level=2)
    
    flows = [
        "Phase 1: Ingestion & Preprocessing\n1. Raw Input: The user records 3-5 seconds of sustained phonation (/a/ vowel) via the browser's Web Audio API at a 44.1 kHz sample rate.\n2. Transmission: The raw .wav blob is sent to the FastAPI backend via a multipart form-data POST request.\n3. Normalization: The backend uses librosa to resample the audio to 22,050 Hz, normalize the amplitude, and remove ambient DC offset noise.",
        "Phase 2: Deterministic Feature Extraction (DSP Engine)\n1. Fundamental Extraction: Praat-Parselmouth extracts the fundamental frequency contour (f0).\n2. Perturbation Calculation: Micro-fluctuations in pitch (Jitter %, absolute, RAP, PPQ) and amplitude (Shimmer, APQ3, APQ5) are calculated.\n3. Noise/Harmonic Measures: Harmonics-to-Noise Ratio (HNR) and Noise-to-Harmonics Ratio (NHR) are derived via autocorrelation.\n4. Non-Linear Dynamics: Pitch Period Entropy (PPE), Detrended Fluctuation Analysis (DFA), and Recurrence Period Density Entropy (RPDE) are calculated to measure the fractal scaling and acoustic chaos indicative of vocal fold rigidity.\nOutput: A 1x22 numerical vector representing the patient's acoustic biomarker profile.",
        "Phase 3: Model Inference Pipeline\n1. Scaling: The 22-dimensional feature vector is normalized using a pre-fitted Scikit-Learn StandardScaler derived from the Oxford training distribution.\n2. Parallel Execution: The scaled vector is passed simultaneously to:\n   - SVM Engine: Generates a probability distribution using an RBF kernel.\n   - Random Forest Engine: Generates a probability using ensemble decision trees.\n   - CNN-BiLSTM Engine: Identifies complex non-linear spatial dependencies in the feature array.\n   - Autoencoder Engine: Attempts to reconstruct the vector. The Mean Squared Error (MSE) is used as an unsupervised anomaly detection z-score.\n   - Telemonitoring Engine: A regression model estimates the patient's UPDRS scores.",
        "Phase 4: Synthesis & Output\n1. The backend averages the probabilities from the classification models to create an ensemble_pd_probability.\n2. The results, UPDRS scores, anomaly metrics, and Z-score comparisons against the normative datasets are serialized into a JSON payload.\n3. The frontend parses the JSON, dynamically updating the Chart.js radar charts, the SVG diagnostic gauge, and rendering the final clinical verdict to the user."
    ]
    for flow in flows:
        p = doc.add_paragraph()
        split_flow = flow.split("\n", 1)
        p.add_run(split_flow[0]).bold = True
        if len(split_flow) > 1:
            p.add_run("\n" + split_flow[1])

    output_path = r"C:\Users\Raaghav Agarwal\.gemini\antigravity-ide\scratch\parkinsons-voice-companion\Parkinsons_Voice_Companion_Report.docx"
    doc.save(output_path)
    print(f"Document saved to {output_path}")

if __name__ == '__main__':
    create_word_doc()
